"""
SMART EVENT PLANNING PLATFORM - Milestone 1
Features:
- Event Creation
- Venue Onboarding & Assignment
- Resource Onboarding & Allocation
- Event Report
- Excel Persistence
"""

import os
import openpyxl


class Resource:
    def __init__(self, resource_id, name, total_quantity):
        self.resource_id = int(resource_id)
        self.name = name
        self.total_quantity = int(total_quantity)


class Venue:
    def __init__(self, venue_id, name, capacity="Not Set"):
        self.venue_id = int(venue_id)
        self.name = name
        self.capacity = capacity


class Event:
    def __init__(self, event_id, name, date="2026-10-10", start_time="10:00 AM", end_time="05:00 PM", assigned_venue_ids=None, allocated_resources=None):
        self.id = int(event_id)         
        self.event_id = int(event_id)  
        self.name = name
        self.date = date
        self.start_time = start_time
        self.end_time = end_time
        self.assigned_venue_ids = assigned_venue_ids if assigned_venue_ids is not None else []
        self.allocated_resources = allocated_resources if allocated_resources is not None else {}
        self.attendees = []            
        self.vendors = []               


events = []
venues = []
resources = []
next_event_id = 1

EXCEL_FILE = "Smart_Event_Platform.xlsx"


def save_m1_to_excel():
    if os.path.exists(EXCEL_FILE):
        wb = openpyxl.load_workbook(EXCEL_FILE)
    else:
        wb = openpyxl.Workbook()

    if "Events" in wb.sheetnames:
        ws_e = wb["Events"]
        ws_e.delete_rows(1, ws_e.max_row)
    else:
        ws_e = wb.create_sheet("Events")

    ws_e.append(["Event ID", "Event Name", "Date", "Start Time", "End Time", "Assigned Venue IDs", "Allocated Resources"])
    for e in events:
        v_str = ",".join(map(str, e.assigned_venue_ids))
        r_str = ";".join([f"{r_id}:{qty}" for r_id, qty in e.allocated_resources.items()])
        ws_e.append([e.id, e.name, e.date, e.start_time, e.end_time, v_str, r_str])

    if "Venues" in wb.sheetnames:
        ws_v = wb["Venues"]
        ws_v.delete_rows(1, ws_v.max_row)
    else:
        ws_v = wb.create_sheet("Venues")

    ws_v.append(["Venue ID", "Venue Name", "Capacity"])
    for v in venues:
        ws_v.append([v.venue_id, v.name, str(v.capacity)])

    if "Resources" in wb.sheetnames:
        ws_r = wb["Resources"]
        ws_r.delete_rows(1, ws_r.max_row)
    else:
        ws_r = wb.create_sheet("Resources")

    ws_r.append(["Resource ID", "Resource Name", "Total Quantity"])
    for r in resources:
        ws_r.append([r.resource_id, r.name, r.total_quantity])

    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])

    wb.save(EXCEL_FILE)


def load_m1_from_excel():
    global next_event_id
    if not os.path.exists(EXCEL_FILE):
        # Default starter events if file doesn't exist
        events.append(Event(1, "AI Workshop", "2026-10-10"))
        events.append(Event(2, "Python Bootcamp", "2026-10-15"))
        next_event_id = 3
        return

    wb = openpyxl.load_workbook(EXCEL_FILE)

    if "Venues" in wb.sheetnames:
        ws_v = wb["Venues"]
        for row in list(ws_v.iter_rows(values_only=True))[1:]:
            if row and row[0] is not None:
                venues.append(Venue(row[0], row[1], row[2]))

    if "Resources" in wb.sheetnames:
        ws_r = wb["Resources"]
        for row in list(ws_r.iter_rows(values_only=True))[1:]:
            if row and row[0] is not None:
                resources.append(Resource(row[0], row[1], row[2]))

    if "Events" in wb.sheetnames:
        ws_e = wb["Events"]
        max_id = 0
        for row in list(ws_e.iter_rows(values_only=True))[1:]:
            if row and row[0] is not None:
                v_ids = [int(v) for v in str(row[5]).split(",") if str(v).strip()]
                r_alloc = {}
                if row[6] and str(row[6]).strip():
                    for item in str(row[6]).split(";"):
                        if ":" in item:
                            r_id, qty = item.split(":")
                            r_alloc[int(r_id)] = int(qty)
                e = Event(row[0], row[1], row[2], row[3], row[4], v_ids, r_alloc)
                events.append(e)
                if e.id > max_id:
                    max_id = e.id
        next_event_id = max_id + 1

    if not events:
        events.append(Event(1, "AI Workshop", "2026-10-10"))
        events.append(Event(2, "Python Bootcamp", "2026-10-15"))
        next_event_id = 3


load_m1_from_excel()


def find_event(event_id):
    for event in events:
        if event.id == int(event_id):
            return event
    return None


def find_venue(venue_id):
    for venue in venues:
        if venue.venue_id == int(venue_id):
            return venue
    return None


def find_resource(resource_id):
    for resource in resources:
        if resource.resource_id == int(resource_id):
            return resource
    return None


def create_event():
    global next_event_id
    print("\n--- Create Event ---")
    name = input("Event Name: ").strip()
    if not name:
        print("Event Name cannot be empty!")
        return

    date = input("Date (YYYY-MM-DD): ").strip()
    start_time_str = input("Start Time (e.g., 10:00 AM): ").strip()
    end_time_str = input("End Time (e.g., 05:00 PM): ").strip()

    new_event = Event(next_event_id, name, date, start_time_str, end_time_str)
    events.append(new_event)
    save_m1_to_excel()
    print(f"Event Created Successfully! Event ID: {next_event_id}")
    next_event_id += 1


def add_venue():
    print("\n--- Add Venue ---")
    try:
        venue_id = int(input("Enter Venue ID: "))
    except ValueError:
        print("Invalid Venue ID!")
        return

    if find_venue(venue_id) is not None:
        print(f"Venue ID {venue_id} already exists!")
        return

    name = input("Venue Name: ").strip()
    capacity = input("Capacity (Press Enter for 'Not Set'): ").strip() or "Not Set"

    venues.append(Venue(venue_id, name, capacity))
    save_m1_to_excel()
    print("Venue Added Successfully!")


def assign_venues_to_event():
    print("\n--- Assign Venues to Event ---")
    try:
        event_id = int(input("Enter Event ID: "))
    except ValueError:
        print("Invalid Event ID!")
        return

    event = find_event(event_id)
    if not event:
        print("Event Not Found!")
        return

    raw_input = input("Enter Venue IDs (comma-separated, e.g., 101, 102): ").strip()
    venue_ids = [int(x.strip()) for x in raw_input.split(",") if x.strip() and x.strip().isdigit()]

    event.assigned_venue_ids = venue_ids
    save_m1_to_excel()
    print("Venues Assigned Successfully!")


def add_resource():
    print("\n--- Add Resource ---")
    try:
        resource_id = int(input("Enter Resource ID: "))
        qty = int(input("Quantity: "))
    except ValueError:
        print("Invalid Input!")
        return

    name = input("Resource Name: ").strip()
    resources.append(Resource(resource_id, name, qty))
    save_m1_to_excel()
    print("Resource Added Successfully!")


def allocate_resource_to_event():
    print("\n--- Allocate Resource to Event ---")
    try:
        event_id = int(input("Enter Event ID: "))
        resource_id = int(input("Enter Resource ID: "))
        req_qty = int(input("Quantity: "))
    except ValueError:
        print("Invalid Input!")
        return

    event = find_event(event_id)
    res = find_resource(resource_id)
    if not event or not res:
        print("Event or Resource Not Found!")
        return

    event.allocated_resources[resource_id] = req_qty
    save_m1_to_excel()
    print("Resource Allocated Successfully!")


def event_report():
    print("\n--- EVENT REPORT ---")
    for e in events:
        print(f"\nID: {e.id} | Name: {e.name} | Date: {e.date}")


def main():
    while True:
        print("""
=====================================================
 SMART EVENT PLANNING PLATFORM - MILESTONE 1
=====================================================
1. Create Event
2. Add Venue
3. Assign Venues to Event
4. Add Resource
5. Allocate Resource to Event
6. View Event Report
7. Exit
=====================================================
""")
        choice = input("Enter Choice (1-7): ").strip()
        if choice == "1": create_event()
        elif choice == "2": add_venue()
        elif choice == "3": assign_venues_to_event()
        elif choice == "4": add_resource()
        elif choice == "5": allocate_resource_to_event()
        elif choice == "6": event_report()
        elif choice == "7": break


if __name__ == "__main__":
    main()