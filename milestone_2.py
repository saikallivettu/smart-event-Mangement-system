"""
SMART EVENT PLANNING PLATFORM WITH RESOURCE COORDINATION SYSTEM - Milestone 2
Features:
- Attendee Registration & Verification
- Gmail OTP Email Verification
- Ticket Generation & Real QR Code Creation
- QR Code Email Delivery
- Live OpenCV Camera QR Attendance Scanner
- Vendor Onboarding & Event Assignment
- Consolidated M1 + M2 Reporting & Excel Sync
"""

import os
import random
import smtplib
from email.message import EmailMessage
import openpyxl
import qrcode
import cv2

try:
    from milestone_1 import (
        events, venues, resources, EXCEL_FILE,
        save_m1_to_excel, find_event, find_venue, find_resource
    )
except ImportError:
    print("Error: 'milestone_1.py' not found! Make sure both milestone files are in the same directory.")
    exit(1)


SENDER_EMAIL = "kallivettup@gmail.com"  # Gmail address
APP_PASSWORD = "bshv souh lkth yubc"    # your 16-character App Password here


class Attendee:
    def __init__(self, reg_id, name, email, phone, ticket_id, status="Registered"):
        self.reg_id = reg_id
        self.name = name
        self.email = email
        self.phone = phone
        self.ticket_id = ticket_id
        self.status = status


class Vendor:
    def __init__(self, vendor_id, name, service):
        self.id = int(vendor_id)
        self.name = name
        self.service = service


vendors = []
ticket_number = 1001

for e in events:
    if not hasattr(e, 'attendees'):
        e.attendees = []
    if not hasattr(e, 'vendors'):
        e.vendors = []


def save_m2_to_excel():
    save_m1_to_excel()

    if os.path.exists(EXCEL_FILE):
        wb = openpyxl.load_workbook(EXCEL_FILE)
    else:
        wb = openpyxl.Workbook()

    if "Attendees" in wb.sheetnames:
        ws_a = wb["Attendees"]
        ws_a.delete_rows(1, ws_a.max_row)
    else:
        ws_a = wb.create_sheet("Attendees")

    ws_a.append(["Registration ID", "Name", "Email", "Phone", "Event ID", "Ticket ID", "Status"])
    for e in events:
        for a in getattr(e, 'attendees', []):
            ws_a.append([a.reg_id, a.name, a.email, a.phone, e.event_id, a.ticket_id, a.status])

    if "Vendors" in wb.sheetnames:
        ws_v = wb["Vendors"]
        ws_v.delete_rows(1, ws_v.max_row)
    else:
        ws_v = wb.create_sheet("Vendors")

    ws_v.append(["Vendor ID", "Vendor Name", "Service"])
    for v in vendors:
        ws_v.append([v.id, v.name, v.service]
                    
    if "Vendor Assignments" in wb.sheetnames:
        ws_va = wb["Vendor Assignments"]
        ws_va.delete_rows(1, ws_va.max_row)
    else:
        ws_va = wb.create_sheet("Vendor Assignments")

    ws_va.append(["Event ID", "Vendor ID", "Service"])
    for e in events:
        for v in getattr(e, 'vendors', []):
            ws_va.append([e.event_id, v.id, v.service])

    wb.save(EXCEL_FILE)


def load_m2_from_excel():
    global ticket_number
    if not os.path.exists(EXCEL_FILE):
        return

    wb = openpyxl.load_workbook(EXCEL_FILE)

    if "Vendors" in wb.sheetnames:
        ws_v = wb["Vendors"]
        for row in list(ws_v.iter_rows(values_only=True))[1:]:
            if row and row[0] is not None:
                if find_vendor(int(row[0])) is None:
                    vendors.append(Vendor(row[0], row[1], row[2]))

    if "Attendees" in wb.sheetnames:
        ws_a = wb["Attendees"]
        max_tkt = 1000
        for row in list(ws_a.iter_rows(values_only=True))[1:]:
            if row and row[0] is not None:
                e_id = int(row[4])
                event = find_event(e_id)
                if event:
                    if not hasattr(event, 'attendees'):
                        event.attendees = []
                    attendee = Attendee(row[0], row[1], row[2], row[3], row[5], row[6])
                    event.attendees.append(attendee)

                tkt_str = str(row[5]).replace("TKT", "")
                if tkt_str.isdigit():
                    num = int(tkt_str)
                    if num > max_tkt:
                        max_tkt = num
        ticket_number = max_tkt + 1

    if "Vendor Assignments" in wb.sheetnames:
        ws_va = wb["Vendor Assignments"]
        for row in list(ws_va.iter_rows(values_only=True))[1:]:
            if row and row[0] is not None:
                event = find_event(int(row[0]))
                vendor = find_vendor(int(row[1]))
                if event and vendor:
                    if not hasattr(event, 'vendors'):
                        event.vendors = []
                    if vendor not in event.vendors:
                        event.vendors.append(vendor)


def find_vendor(vendor_id):
    for v in vendors:
        if v.id == int(vendor_id):
            return v
    return None


load_m2_from_excel()



def send_otp(email):
    otp = str(random.randint(100000, 999999))
    msg = EmailMessage()
    msg["Subject"] = "Smart Event Platform - OTP Verification"
    msg["From"] = SENDER_EMAIL
    msg["To"] = email
    msg.set_content(
        f"Hello,\n\nYour OTP for Smart Event Platform registration is:\n\n{otp}\n\n"
        f"Please do not share this OTP with anyone.\n\nThank you."
    )

    try:
    
        with smtplib.SMTP("smtp.gmail.com", 587) as server:
            server.starttls()
            server.login(SENDER_EMAIL, APP_PASSWORD)
            server.send_message(msg)
        print("\nOTP sent successfully to", email)
        return otp
    except Exception as e:
        print("\nFailed to send OTP email.")
        print("SMTP Error Info:", e)
        print(f"(Console Backup OTP for testing: {otp})")
        return otp


def verify_otp(email):
    otp = send_otp(email)
    if otp is None:
        return False

    for attempt in range(3):
        entered_otp = input("Enter OTP: ").strip()
        if entered_otp == otp:
            print("Email verification successful.")
            return True
        print("Wrong OTP.")

    print("Maximum OTP attempts reached.")
    return False

def generate_qr(ticket_id, reg_id, event_id):
    data = f"{ticket_id}|{reg_id}|{event_id}"
    qr = qrcode.make(data)
    filename = f"{ticket_id}.png"
    qr.save(filename)
    return filename



def send_ticket_email(email, name, ticket, qr_file):
    msg = EmailMessage()
    msg["Subject"] = "Your Smart Event Ticket Pass"
    msg["From"] = SENDER_EMAIL
    msg["To"] = email
    msg.set_content(
        f"Hello {name},\n\nYour event registration was successful.\n\n"
        f"Ticket ID: {ticket}\n\nYour QR code is attached to this email.\n"
        f"Please present this QR code at the event entrance.\n\nThank you."
    )

    try:
        with open(qr_file, "rb") as f:
            msg.add_attachment(
                f.read(),
                maintype="image",
                subtype="png",
                filename=qr_file
            )

        with smtplib.SMTP_SSL("smtp.gmail.com", 465) as server:
            server.login(SENDER_EMAIL, APP_PASSWORD)
            server.send_message(msg)

        print("Ticket and QR code emailed to", email)
        return True
    except Exception as e:
        print("Ticket email error:", e)
        return False



def view_events():
    print("\n--- Available Events ---")
    for e in events:
        print(f"ID: {e.event_id} | Name: {e.name} | Date: {e.date}")


def register_attendee():
    global ticket_number
    if len(events) == 0:
        print("No events available! Create an event in Milestone 1 first.")
        return

    view_events()
    try:
        event = find_event(int(input("Event ID: ")))
    except ValueError:
        print("Invalid Input")
        return

    if event is None:
        print("Event Not Found")
        return

    if not hasattr(event, 'attendees'):
        event.attendees = []

    reg = input("Registration ID: ").strip()
    name = input("Name: ").strip()
    email = input("Email: ").strip()
    phone = input("Phone: ").strip()

    for a in event.attendees:
        if a.email.lower() == email.lower():
            print("Already Registered for this event!")
            return

    print("\n========== EMAIL VERIFICATION ==========")
    if not verify_otp(email):
        print("Registration cancelled.")
        return

    ticket = f"TKT{ticket_number}"
    ticket_number += 1

    attendee = Attendee(reg, name, email, phone, ticket)
    event.attendees.append(attendee)

    qr_file = generate_qr(ticket, reg, event.event_id)
    send_ticket_email(email, name, ticket, qr_file)

    save_m2_to_excel()

    print("\n========== REGISTRATION SUCCESSFUL ==========")
    print("Registration ID:", reg)
    print("Name           :", name)
    print("Email          :", email)
    print("Phone          :", phone)
    print("Ticket ID      :", ticket)
    print("QR Image File  :", qr_file)


def view_attendees():
    view_events()
    try:
        event = find_event(int(input("Event ID: ")))
    except ValueError:
        print("Invalid Input")
        return

    if event is None:
        print("Event Not Found")
        return

    attendees_list = getattr(event, 'attendees', [])
    if not attendees_list:
        print("No Registrations found for this event.")
        return

    print(f"\nRegistered Participants for '{event.name}':")
    for a in attendees_list:
        print(f"Name: {a.name:<15} | Ticket: {a.ticket_id} | Status: {a.status}")


def mark_attendance():
    view_events()
    try:
        event = find_event(int(input("Event ID: ")))
    except ValueError:
        print("Invalid Input")
        return

    if event is None:
        print("Event Not Found")
        return

    print("\n====================================")
    print("  QR ATTENDANCE SCANNER (OpenCV)")
    print("  Show the participant's QR code")
    print("  Press 'q' in camera view to exit")
    print("====================================")

    camera = cv2.VideoCapture(0)
    if not camera.isOpened():
        print("Camera could not be opened.")
        return

    detector = cv2.QRCodeDetector()

    while True:
        ret, frame = camera.read()
        if not ret:
            print("Camera capture error.")
            break

        data, points, _ = detector.detectAndDecode(frame)
        cv2.imshow("QR Attendance Scanner", frame)

        if data:
            print("\nScanned QR Data:", data)
            parts = data.split("|")

            if len(parts) != 3:
                print("Invalid QR Code Format!")
                continue

            ticket = parts[0]
            reg_id = parts[1]
            qr_event_id = parts[2]

            try:
                qr_event_id = int(qr_event_id)
            except ValueError:
                print("Invalid Event ID inside QR!")
                continue

            if qr_event_id != event.event_id:
                print("QR belongs to a different event!")
                continue

            found = False
            for a in getattr(event, 'attendees', []):
                if a.ticket_id == ticket:
                    found = True
                    if a.reg_id != reg_id:
                        print("Registration verification failed!")
                        break

                    if a.status == "Checked In":
                        print("ALREADY CHECKED IN!")
                        camera.release()
                        cv2.destroyAllWindows()
                        return

                    a.status = "Checked In"
                    save_m2_to_excel()

                    print("\n========== ATTENDANCE SUCCESS ==========")
                    print("Name           :", a.name)
                    print("Registration ID:", a.reg_id)
                    print("Ticket ID      :", a.ticket_id)
                    print("Event          :", event.name)
                    print("Status         :", a.status)

                    camera.release()
                    cv2.destroyAllWindows()
                    return

            if not found:
                print("Invalid Ticket for this event!")

        if cv2.waitKey(1) & 0xFF == ord("q"):
            print("Scanner closed.")
            break

    camera.release()
    cv2.destroyAllWindows()


def add_vendor():
    try:
        vid = int(input("Vendor ID (numeric): "))
    except ValueError:
        print("Invalid Vendor ID!")
        return

    if find_vendor(vid) is not None:
        print("Vendor ID already exists!")
        return

    name = input("Vendor Name: ").strip()
    service = input("Service Provided: ").strip()

    vendors.append(Vendor(vid, name, service))
    save_m2_to_excel()
    print(f"Vendor '{name}' Added Successfully!")


def view_vendors():
    if not vendors:
        print("No Vendors Available.")
        return

    print("\n--- Onboarded Vendors ---")
    for v in vendors:
        print(f"ID: {v.id} | Name: {v.name} | Service: {v.service}")


def assign_vendor():
    view_events()
    try:
        event = find_event(int(input("Event ID: ")))
    except ValueError:
        print("Invalid Input")
        return

    if event is None:
        print("Event Not Found")
        return

    view_vendors()
    if not vendors:
        return

    try:
        vendor = find_vendor(int(input("Vendor ID to assign: ")))
    except ValueError:
        print("Invalid Input")
        return

    if vendor is None:
        print("Vendor Not Found")
        return

    if not hasattr(event, 'vendors'):
        event.vendors = []

    if vendor in event.vendors:
        print("Vendor is already assigned to this event.")
        return

    event.vendors.append(vendor)
    save_m2_to_excel()
    print(f"Vendor '{vendor.name}' Assigned to '{event.name}' Successfully!")


def report():
    print("\n================ SYSTEM REPORT (M1 + M2) ================")
    if len(events) == 0:
        print("No Events Found.")
        return

    for e in events:
        attendees_list = getattr(e, 'attendees', [])
        vendors_list = getattr(e, 'vendors', [])

        checked = sum(1 for a in attendees_list if a.status == "Checked In")

        print(f"\nEvent Name   : {e.name} (ID: {e.event_id})")
        print(f"Date & Time  : {e.date} | {e.start_time} - {e.end_time}")
        print(f"Registrations: {len(attendees_list)}")
        print(f"Checked In   : {checked}")

        print("Vendors Assigned:")
        if not vendors_list:
            print("  None")
        else:
            for v in vendors_list:
                print(f"  - {v.name} ({v.service})")
    print("========================================================")


def main():
    while True:
        print("""
=============================================================
 SMART EVENT PLANNING PLATFORM - MILESTONE 2
=============================================================
1. Register Attendee
2. View Attendees
3. Mark Attendance using QR Scanner (OpenCV)
4. Add Vendor
5. View Vendors
6. Assign Vendor to Event
7. Report
8. Exit
=============================================================
""")
        ch = input("Choice (1-8): ").strip()

        if ch == "1":
            register_attendee()
        elif ch == "2":
            view_attendees()
        elif ch == "3":
            mark_attendance()
        elif ch == "4":
            add_vendor()
        elif ch == "5":
            view_vendors()
        elif ch == "6":
            assign_vendor()
        elif ch == "7":
            report()
        elif ch == "8":
            print("Exiting Milestone 2 System. Thank You!")
            break
        else:
            print("Invalid Choice! Please enter a number from 1 to 8.")


if __name__ == "__main__":
    main()